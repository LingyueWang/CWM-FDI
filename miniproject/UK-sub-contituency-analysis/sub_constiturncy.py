#sub_consitituency digital equity index calculation and analysis
#Lingyue Wang 12/06/2026

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook


BASE_DIR = Path(__file__).resolve().parent

INPUT_FILE = BASE_DIR / "normalized_subconstituency_data.xlsx"
WEIGHTS_WITH_SPEED_FILE = BASE_DIR / "weights_with_speed.json"
WEIGHTS_WITHOUT_SPEED_FILE = BASE_DIR / "weights_without_speed.json"

SHEET_NAME = "Sheet1"
HEADER_SCAN_ROWS = 20

ID_HEADERS = ["Area code", "Area name"]

METRICS_WITH_SPEED = [
    "Superfast availability",
    "Gigabit availability",
    "Receiving over 30 Mbps",
    "Below universal service obligation",
    "Receiving under 10 Mbps",
    "Average download speed (Mbps)",
]

METRICS_WITHOUT_SPEED = [
    "Superfast availability",
    "Gigabit availability",
    "Receiving over 30 Mbps",
    "Below universal service obligation",
    "Receiving under 10 Mbps",
]

ALIASES = {
    "Area code": ["Area code"],
    "Area name": ["Area name"],
    "Superfast availability": ["Superfast availability"],
    "Gigabit availability": ["Gigabit availability"],
    "Receiving over 30 Mbps": ["Receiving over 30 Mbps", "Receiving over 30Mbps"],
    "Below universal service obligation": ["Below universal service obligation"],
    "Receiving under 10 Mbps": ["Receiving under 10 Mbps", "Receiving under 10Mbps"],
    "Average download speed (Mbps)": ["Average download speed (Mbps)", "Average download speed"],
}


def normalize_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def load_weights(path: Path, metrics: list[str]) -> dict[str, float]:
    with path.open("r", encoding="utf-8") as f:
        raw = json.load(f)

    weights = {str(k).strip(): float(v) for k, v in raw.items()}

    missing = [m for m in metrics if m not in weights]
    if missing:
        raise ValueError(f"Missing weights in {path.name} for: {missing}")

    total = sum(weights[m] for m in metrics)
    if abs(total - 1.0) > 1e-6:
        raise ValueError(f"Weights in {path.name} must sum to 1. Current sum = {total}")

    return weights


def choose_sheet(workbook):
    if SHEET_NAME in workbook.sheetnames:
        return workbook[SHEET_NAME]
    return workbook[workbook.sheetnames[0]]


def find_header_row_and_columns(ws, required_headers: list[str]) -> tuple[int, dict[str, int]]:
    for row_num in range(1, HEADER_SCAN_ROWS + 1):
        row = ws[row_num]
        header_map: dict[str, int] = {}

        for cell in row:
            if cell.value is not None:
                header_map[normalize_key(cell.value)] = cell.column

        resolved_columns: dict[str, int] = {}
        ok = True

        for header in required_headers:
            found_col = None
            for alias in ALIASES.get(header, [header]):
                key = normalize_key(alias)
                if key in header_map:
                    found_col = header_map[key]
                    break
            if found_col is None:
                ok = False
                break
            resolved_columns[header] = found_col

        if ok:
            return row_num, resolved_columns

    raise ValueError("Could not find a header row with the required columns.")


def write_workbook(path: Path, rows: list[list[object]]) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    for row in rows:
        ws.append(row)
    wb.save(path)


def process_mode(mode_name: str, metrics: list[str], weights_path: Path) -> None:
    weights = load_weights(weights_path, metrics)

    weighted_output_file = BASE_DIR / f"weighted_parameters_{mode_name}.xlsx"
    index_output_file = BASE_DIR / f"digital_equity_index_{mode_name}.xlsx"

    wb = load_workbook(INPUT_FILE, data_only=True, read_only=True)
    ws = choose_sheet(wb)

    required_headers = ID_HEADERS + metrics
    header_row, cols = find_header_row_and_columns(ws, required_headers)

    weighted_rows = [["Area code", "Area name"]]
    weighted_rows[0] += [f"Normalized - {m}" for m in metrics]
    weighted_rows[0] += [f"Weighted - {m}" for m in metrics]
    weighted_rows[0] += ["Digital Equity Index"]

    index_rows = [["Area code", "Area name", "Digital Equity Index"]]

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        def get_cell(col_num: int):
            idx = col_num - 1
            return row[idx] if idx < len(row) else None

        area_code = get_cell(cols["Area code"])
        area_name = get_cell(cols["Area name"])

        if area_code is None and area_name is None:
            continue

        normalized_values = {}
        weighted_values = {}
        has_data = False

        for metric in metrics:
            value = get_cell(cols[metric])
            if value is None:
                raise ValueError(f"Missing value for {metric} in row {area_code}")

            norm = float(value)  # already normalized in the input workbook
            norm = max(0.0, min(1.0, norm))
            weighted = norm * weights[metric]

            normalized_values[metric] = norm
            weighted_values[metric] = weighted
            has_data = True

        if not has_data:
            continue

        index_value = sum(weighted_values[m] for m in metrics)

        weighted_rows.append(
            [area_code, area_name]
            + [round(normalized_values[m], 6) for m in metrics]
            + [round(weighted_values[m], 6) for m in metrics]
            + [round(index_value, 6)]
        )

        index_rows.append([area_code, area_name, round(index_value, 6)])

    write_workbook(weighted_output_file, weighted_rows)
    write_workbook(index_output_file, index_rows)

    print(f"Created: {weighted_output_file}")
    print(f"Created: {index_output_file}")


def main() -> None:
    process_mode("with_speed", METRICS_WITH_SPEED, WEIGHTS_WITH_SPEED_FILE)
    process_mode("without_speed", METRICS_WITHOUT_SPEED, WEIGHTS_WITHOUT_SPEED_FILE)


if __name__ == "__main__":
    main()